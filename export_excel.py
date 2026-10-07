"""Export Excel des offres suivies, régénéré à partir de la base (data/)."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from annonce import FIELDS, PARIS, extra_fields, is_in_zone, parse_date, parse_dt, pick
from stockage import Store

FONT = "Arial"
SECTION_FILLS = {
    "Suivi": "7A0019", "Localisation": "9C1C2E", "Loyer et frais": "B3001B", "Logement": "8A3B12",
    "Énergie": "5E4B10", "Résidence": "2F4B5E", "Conditions": "3B3B6B", "Description": "4A4A4A",
    "Références": "333333", "Autres champs": "5A5A5A",
}
FMT = {
    "datetime": "dd/mm/yyyy hh:mm",
    "date": "dd/mm/yyyy",
    "money": '#,##0.00 "€"',
    "area": "0.0#",
    "int": "0",
}


def _text(value) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))[:32000]


def _value(value, kind: str):
    if value in (None, ""):
        return None
    if kind == "datetime":
        return value.astimezone(PARIS).replace(tzinfo=None) if isinstance(value, datetime) else None
    if kind == "date":
        return value if isinstance(value, date) else parse_date(value)
    if kind in ("money", "area"):
        try:
            return float(value)
        except (TypeError, ValueError):
            return _text(value)
    if kind == "int":
        try:
            return int(value)
        except (TypeError, ValueError):
            return _text(value)
    if kind == "bool":
        return "Oui" if value else "Non"
    if isinstance(value, bool):
        return "Oui" if value else "Non"
    if isinstance(value, (int, float)):
        return value
    return _text(value)


def _sort_key(rec: dict):
    return (rec.get("premiere_detection") or "", pick(rec, "startDateOfPublication") or "", rec["guid"])


def _header(ws, columns: list[tuple[str, str, int]]) -> None:
    for i, (label, section, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=label)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=SECTION_FILLS.get(section, "333333"))
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "B2"


def _offers_sheet(ws, records: list[dict], extras: list[str]) -> None:
    columns = [(f.label, f.section, f.width) for f in FIELDS] + [(f"Autre · {p}", "Autres champs", 18) for p in extras]
    _header(ws, columns)
    col = {f.label: i for i, f in enumerate(FIELDS, start=1)}
    c_pub = get_column_letter(col["Publiée le (site)"])
    c_det = get_column_letter(col["Mise en ligne détectée le"])
    c_ret = get_column_letter(col["Retrait détecté le"])
    link_font = Font(name=FONT, size=10, color="0563C1", underline="single")

    for r, rec in enumerate(records, start=2):
        ws.row_dimensions[r].height = 15  # lignes compactes malgré les descriptions sur plusieurs lignes
        for c, f in enumerate(FIELDS, start=1):
            cell = ws.cell(row=r, column=c)
            if f.kind == "formula":  # durée en ligne : depuis la publication (ou la détection) jusqu'au retrait ou à aujourd'hui
                cell.value = (f'=IF(AND({c_pub}{r}="",{c_det}{r}=""),"",'
                              f'INT(IF({c_ret}{r}="",NOW(),{c_ret}{r})-IF({c_pub}{r}="",{c_det}{r},{c_pub}{r})))')
                cell.number_format = "0"
                continue
            raw = f.get(rec)
            if f.kind == "link":
                cell.value = "Voir l'annonce"
                cell.hyperlink = raw
                cell.font = link_font
                continue
            cell.value = _value(raw, f.kind)
            if f.kind in FMT and cell.value is not None:
                cell.number_format = FMT[f.kind]
        if extras:
            ex = extra_fields(rec)
            for j, path in enumerate(extras, start=len(FIELDS) + 1):
                if path in ex:
                    ws.cell(row=r, column=j, value=_value(ex[path], "text"))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(2, len(records) + 1)}"


def _history_sheet(ws, events: list[dict]) -> None:
    columns = [("Date et heure", "Suivi", 17), ("Événement", "Suivi", 15), ("Ville", "Localisation", 26),
               ("Code postal", "Localisation", 10), ("Typologie", "Logement", 9), ("Loyer", "Loyer et frais", 11),
               ("Détails", "Description", 70), ("Référence", "Références", 11), ("Lien", "Références", 12)]
    _header(ws, columns)
    link_font = Font(name=FONT, size=10, color="0563C1", underline="single")
    for r, ev in enumerate(reversed(events), start=2):
        when = parse_dt(ev.get("horodatage"))
        ws.cell(row=r, column=1, value=when.replace(tzinfo=None) if when else ev.get("horodatage"))
        ws.cell(row=r, column=1).number_format = FMT["datetime"]
        for c, key in enumerate(["evenement", "ville", "code_postal", "typologie", "loyer", "details", "reference"],
                                start=2):
            ws.cell(row=r, column=c, value=_text(ev.get(key, "")) or None)
        if ev.get("guid"):
            cell = ws.cell(row=r, column=9, value="Voir")
            cell.hyperlink = config.DETAIL_URL.format(guid=ev["guid"])
            cell.font = link_font
    ws.auto_filter.ref = f"A1:I{max(2, len(events) + 1)}"


def _about_sheet(ws, store: Store, n_offers: int, n_zone: int, n_history: int) -> None:
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 90
    bold = Font(name=FONT, bold=True, size=10)
    title = ws.cell(row=1, column=1, value="Suivi des offres — logement-actionlogement.fr")
    title.font = Font(name=FONT, bold=True, size=14)
    generated = datetime.now(PARIS).replace(tzinfo=None, microsecond=0)
    debut = parse_dt(store.state.get("debut_suivi"))
    last = parse_dt(store.state.get("dernier_passage_complet"))
    rows = [
        ("Fichier généré le", generated, FMT["datetime"]),
        ("Suivi démarré le", debut.replace(tzinfo=None) if debut else None, FMT["datetime"]),
        ("Dernière vérification complète", last.replace(tzinfo=None) if last else None, FMT["datetime"]),
        ("Offres enregistrées", f"=COUNTA(Offres!A2:A{n_offers + 1})", "0"),
        ("Dont en ligne", f'=COUNTIF(Offres!A2:A{n_offers + 1},"En ligne")', "0"),
        ("Dont retirées", f'=COUNTIF(Offres!A2:A{n_offers + 1},"Retirée")', "0"),
        ("Offres dans la zone d'alerte", f"=COUNTA('Zone d''alerte'!A2:A{n_zone + 1})", "0"),
        ("Événements dans l'historique", f"=COUNTA(Historique!A2:A{n_history + 1})", "0"),
        (None, None, None),
        ("Comment lire les dates", None, None),
        ("Publiée le (site)", "Date de publication affichée par le site (au jour près).", None),
        ("Mise en ligne détectée le", "Première vérification où l'offre est apparue. Précision : l'intervalle entre deux "
                                      "vérifications (20 min sur GitHub, plus si GitHub est chargé).", None),
        ("Déjà en ligne au démarrage du suivi", "Oui = l'offre existait avant le premier passage : sa date de détection "
                                                "est celle du démarrage, pas sa vraie mise en ligne.", None),
        ("Vue en ligne pour la dernière fois", "Dernière vérification complète où l'offre était encore présente.", None),
        ("Retrait détecté le", f"Première vérification où l'offre a disparu, validée après {config.REMOVAL_CONFIRM_RUNS} "
                               "vérifications complètes d'affilée sans elle (pour éviter les faux retraits).", None),
        ("Durée en ligne (jours)", "Calculée par Excel : de la publication (ou de la détection) jusqu'au retrait, "
                                   "ou jusqu'à aujourd'hui si l'offre est encore en ligne.", None),
        ("Remises en ligne", "Nombre de fois où une offre retirée est réapparue.", None),
        (None, None, None),
        ("Zone d'alerte (codes INSEE)", ", ".join(config.ALERT_ZONES), None),
        ("Colonnes « Autre · … »", "Champs bruts renvoyés par le site qui n'ont pas de colonne dédiée : rien n'est perdu.",
         None),
        ("Source", f"API publique du site ({config.API_BASE}), recherche : {config.SEARCH_BODY}", None),
        ("Données de référence", "Dossier data/ (un fichier JSON par offre + historique.csv). Ce classeur est "
                                 "régénéré à partir de ces fichiers et peut être supprimé sans perte.", None),
    ]
    for r, (label, value, number_format) in enumerate(rows, start=3):
        if label is None:
            continue
        a = ws.cell(row=r, column=1, value=label)
        a.font = bold
        if value is not None:
            b = ws.cell(row=r, column=2, value=value)
            b.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
            if number_format:
                b.number_format = number_format
        if value is None:
            a.font = Font(name=FONT, bold=True, size=11)


def build(store: Store, path: Path) -> Path:
    records = sorted(store.records.values(), key=_sort_key, reverse=True)
    zone = [r for r in records if is_in_zone(r)]
    events = store.read_history()
    extras = sorted({p for rec in records for p in extra_fields(rec)})

    wb = Workbook()
    # police par défaut du classeur
    normal = next(s for s in wb._named_styles if s.name == "Normal")
    normal.font = Font(name=FONT, size=10)

    ws = wb.active
    ws.title = "Offres"
    _offers_sheet(ws, records, extras)
    _offers_sheet(wb.create_sheet("Zone d'alerte"), zone, extras)
    _history_sheet(wb.create_sheet("Historique"), events)
    _about_sheet(wb.create_sheet("À propos"), store, len(records), len(zone), len(events))
    ws.cell(row=1, column=[f.label for f in FIELDS].index("Durée en ligne (jours)") + 1).comment = Comment(
        "Calculé à l'ouverture : de la publication (ou de la détection) jusqu'au retrait ou à aujourd'hui.", "Veille")

    wb.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)
    return path
